import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from admin_doc_builder import (
    set_cell_background, set_cell_margins, add_callout,
    add_styled_table, add_header_footer, add_section_h1,
    add_section_h2, add_section_h3, add_p, add_bullet, add_num_step
)

def build_pure_admin_manual(output_path, include_en=True, include_bn=True):
    doc = Document()
    add_header_footer(doc, "Admin Portal Manual / অ্যাডমিন ব্যবহারের নিয়ম")
    
    # Title Cover Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(18)
    title_p.paragraph_format.space_after = Pt(2)
    r_brand = title_p.add_run("SMART WORK FORCE")
    r_brand.bold = True
    r_brand.font.size = Pt(24)
    r_brand.font.name = "Segoe UI"
    r_brand.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Navy
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    r_sub = sub_p.add_run("Admin Operational Manual: How to Monitor & Manage Field Team (Step-by-Step)\nঅ্যাডমিন অপারেশন ম্যানুয়াল: টিম মনিটরিং ও পরিচালনার সহজ নিয়ম")
    r_sub.font.size = Pt(13)
    r_sub.font.name = "Segoe UI"
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    add_callout(
        doc,
        "এই ম্যানুয়ালে কোনো জটিল সফটওয়্যার কোড বা টেকনিক্যাল ভাষা নেই। আপনি কীভাবে লগইন করবেন, কর্মীদের লাইভ অবস্থান ম্যাপে দেখবেন, কে কোথায় কতক্ষণ ছিল (রুট হিস্ট্রি) দেখবেন, সেলফি হাজিরা চেক করবেন, ছুটির দরখাস্ত অ্যাপ্রুভ করবেন এবং নতুন কর্মী অ্যাড করবেন — তা ধাপে ধাপে বুঝিয়ে দেওয়া হয়েছে।\n"
        "This is a practical action-oriented manual for Administrators. It explains step-by-step how to monitor live tracking, check route history, review selfie attendance, approve leaves, and manage employee accounts without any technical jargon.",
        title="ADMIN HOW-TO MANUAL / অ্যাডমিন ব্যবহার নির্দেশিকা",
        border_hex="1E3A8A",
        bg_hex="F0F7FF"
    )
    
    # -------------------------------------------------------------
    # BANGLA SECTION (FIRST & PRIMARY)
    # -------------------------------------------------------------
    if include_bn:
        add_section_h1(doc, "প্রথম অংশ: অ্যাডমিন ব্যবহারের সহজ নিয়ম (বাংলা সংস্করণ)")
        
        # ১. অ্যাডমিন লগইন
        add_section_h2(doc, "১. অ্যাডমিন প্যানেলে লগইন করার নিয়ম")
        add_num_step(doc, 1, "অ্যাপ ওপেন করুন", "Smart Work Force অ্যাপটি ওপেন করুন।")
        add_num_step(doc, 2, "অ্যাডমিন আইডি ও পাসওয়ার্ড দিন", "আপনার জন্য নির্ধারিত অ্যাডমিন Username এবং Password লিখুন।")
        add_num_step(doc, 3, "লগইন চাপুন", "'LOGIN' বাটনে চাপ দিন। আপনি সরাসরি অ্যাডমিন কমান্ড ড্যাশবোর্ডে প্রবেশ করবেন।")
        add_num_step(doc, 4, "ভাষা বদলানো", "ওপরে ডানপাশে '🌐 বাংলা' বাটনে চাপ দিয়ে পুরো অ্যাপ বাংলায় দেখতে পারবেন।")
        
        # ২. ড্যাশবোর্ড থেকে এক নজরে সব দেখা
        add_section_h2(doc, "২. হোম ড্যাশবোর্ড থেকে পুরো টিমের অবস্থা দেখা")
        add_p(doc, "ড্যাশবোর্ডে ঢোকা মাত্রই স্ক্রিনের কার্ডগুলোতে আজকের সারাদিনের অবস্থা এক নজরে দেখতে পাবেন:")
        add_bullet(doc, "মোট কর্মী (Total Users)", "আপনার প্রতিষ্ঠানে মোট কতজন কর্মী আছে।")
        add_bullet(doc, "বর্তমানে সক্রিয় (Active Tracking)", "এই মুহূর্তে কতজন অফিসারের জিপিএস চালু আছে এবং তারা মাঠে কাজ করছেন।")
        add_bullet(doc, "আজকের উপস্থিতি (Today's Punch In)", "আজকে কয়জন অফিসে বা মাঠে হাজিরা দিয়েছেন।")
        add_bullet(doc, "আজকে অনুপস্থিত (Absent)", "আজকে কারা এখনো হাজিরা দেননি।")
        add_bullet(doc, "দেরিতে আসা (Late)", "নির্ধারিত সময়ের পরে কারা হাজিরা দিয়েছেন।")
        add_bullet(doc, "জরুরি সতর্কতা (Attention Required)", "যদি কোনো কর্মী কাজের সময় জিপিএস বন্ধ করে বা ৩ মিনিটের বেশি নেটওয়ার্ক না থাকে, তবে সাথে সাথে লাল সতর্কতা দেখাবে। আপনি সেখান থেকেই সরাসরি তাকে ফোন দিতে পারবেন।")
        
        # ৩. লাইভ ম্যাপে কর্মীদের দেখা
        add_section_h2(doc, "৩. গুগল ম্যাপে কর্মীদের লাইভ অবস্থান দেখা (Live Map)")
        add_num_step(doc, 1, "ম্যাপ মেনুতে যান", "নিচের মেনু থেকে 'Live Map' বা 'ম্যাপ' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "পিনের রঙ দেখে স্থিতি বুঝুন", "গুগল ম্যাপে কর্মীদের রঙিন পিন দেখতে পাবেন:")
        add_bullet(doc, "🟢 সবুজ পিন (Green)", "অফিসার চলমান (বাইকে বা গাড়িতে চলাচল করছেন)।")
        add_bullet(doc, "🔵 নীল পিন (Blue)", "অফিসার কোনো দোকানে বা কাস্টমারের সাথে মিটিংয়ে দাঁড়িয়ে আছেন।")
        add_bullet(doc, "⚪ ধূসর পিন (Gray)", "অফলাইন (গত ৩ মিনিট ধরে নেটওয়ার্ক বা মোবাইল বন্ধ)।")
        add_bullet(doc, "🔴 লাল পিন (Red)", "কর্মীর মোবাইল থেকে জিপিএস লোকেশন বন্ধ করা হয়েছে।")
        add_num_step(doc, 3, "পিনের ওপর চাপ দিন", "যেকোনো পিনে চাপ দিলে অফিসারের নাম, মোবাইল নম্বর, বর্তমান গতি, ব্যাটারির চার্জ এবং ঠিকানা ভেসে উঠবে।")
        add_num_step(doc, 4, "সরাসরি কল দিন", "তথ্য কার্ডের 'Call' বাটনে চাপ দিয়ে সরাসরি তাকে ফোন করতে পারবেন।")
        
        # ৪. রুট হিস্ট্রি ও সারাদিনের ভ্রমণ দেখা
        add_section_h2(doc, "৪. কর্মীর সারাদিনের চলাচলের রুট দেখা (Route History)")
        add_p(doc, "কোনো কর্মী সারাদিনে কোন কোন রাস্তা দিয়ে ঘুরেছে এবং কোথায় কতক্ষণ দাঁড়িয়ে ছিল তা দেখার নিয়ম:")
        add_num_step(doc, 1, "Route History মেনুতে যান", "মেনু থেকে 'Route History' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "কর্মীর নাম ও তারিখ বাছুন", "যে কর্মীর রুট দেখতে চান তার নাম সিলেক্ট করুন এবং ক্যালেন্ডার থেকে তারিখ বেছে নিন।")
        add_num_step(doc, 3, "ভ্রমণ সারাংশ দেখুন", "স্ক্রিনের ওপরে দেখতে পাবেন: মোট কত কিলোমিটার পথ ঘুরেছেন, সকাল কয়টায় শুরু এবং সন্ধ্যায় কয়টায় শেষ করেছেন।")
        add_num_step(doc, 4, "স্টপ ও বিরতি চেক করুন", "ম্যাপে ১, ২, ৩ নম্বর দিয়ে চিহ্নিত স্টপ পিন দেখতে পাবেন। যেকোনো স্টপ পিনে চাপ দিলে সে ওই দোকানে কয়টায় ঢুকেছিল এবং কত মিনিট অবস্থান করেছিল তা স্পষ্ট দেখতে পাবেন।")
        add_num_step(doc, 5, "ভিডিও প্লেব্যাক চালান", "নিচে থাকা 'Play' বাটনে চাপ দিলে ভিডিওর মতো দেখতে পাবেন সারাদিনে গাড়ি বা বাইক কীভাবে চলেছে।")
        
        # ৫. হাজিরা ও সেলফি চেক করা
        add_section_h2(doc, "৫. কর্মীদের হাজিরা ও সেলফি ছবি চেক করা (Attendance)")
        add_num_step(doc, 1, "Attendance মেনুতে যান", "মেনু থেকে 'Attendance' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "তালিকা দেখুন", "আজকের বা পেছনের যেকোনো তারিখের সকল কর্মীর উপস্থিতির তালিকা দেখতে পাবেন।")
        add_num_step(doc, 3, "সেলফি ছবি বড় করে দেখুন", "কর্মীর নামের পাশে ছোট ছবির ওপর চাপ দিলে তার তোলা আসল সেলফি বড় হয়ে স্ক্রিনে আসবে।")
        add_num_step(doc, 4, "অফিস নাকি ফিল্ড হাজিরা", "সে অফিসের ভেতর থেকে হাজিরা দিয়েছে নাকি অফিসের বাইরে থেকে, তা দেখে ভেরিফাই করতে পারবেন।")
        
        # ৬. কাস্টমার ভিজিট তদারকি
        add_section_h2(doc, "৬. কর্মীদের কাস্টমার ভিজিট রিপোর্ট দেখা")
        add_num_step(doc, 1, "Customer Visits এ যান", "মেনু থেকে 'Visits' বা 'Customer Visits' এ চাপ দিন।")
        add_num_step(doc, 2, "মিটিং রিপোর্ট দেখুন", "মাঠের কর্মীরা সারাদিনে কোন কোন দোকানে গেছে, দোকানদারের নাম, মোবাইল এবং কী কথা হয়েছে (নোট) দেখতে পাবেন।")
        add_num_step(doc, 3, "দোকানের ছবি দেখুন", "অফিসার যদি কাস্টমারের দোকানের ছবি তুলে থাকে, তা জিপিএস লোকেশনসহ দেখতে পাবেন।")
        
        # ৭. ছুটির আবেদন অ্যাপ্রুভ করা
        add_section_h2(doc, "৭. কর্মীদের ছুটির আবেদন অনুমোদন বা বাতিল করা")
        add_num_step(doc, 1, "Leave মেনুতে যান", "মেনু থেকে 'Leave' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "পেন্ডিং তালিকা দেখুন", "কে ছুটির জন্য আবেদন করেছে, ছুটির ধরন (নৈমিত্তিক/অসুস্থতা), তারিখ এবং কারণ দেখতে পাবেন।")
        add_num_step(doc, 3, "অনুমোদন বা বাতিল করুন", "ছুটি মঞ্জুর করতে 'Approve' বাটনে চাপ দিন, অথবা নামঞ্জুর করতে 'Reject' বাটনে চাপ দিন। সাথে সাথে কর্মীর মোবাইলে নোটিফিকেশন পৌঁছে যাবে।")
        
        # ৮. নতুন কর্মী অ্যাড করা
        add_section_h2(doc, "৮. নতুন কর্মীর আইডি তৈরি করা (Create User)")
        add_num_step(doc, 1, "Users মেনুতে যান", "মেনু থেকে 'User Management' বা 'Users' এ চাপ দিন।")
        add_num_step(doc, 2, "+ বাটন চাপুন", "উপরে থাকা '+ Create User' বাটনে চাপ দিন।")
        add_num_step(doc, 3, "তথ্য পূরণ করুন", "কর্মীর পুরো নাম, মোবাইল নম্বর, ইউজারনেম এবং পাসওয়ার্ড লিখুন।")
        add_num_step(doc, 4, "পদবি ও ব্রাঞ্চ বাছুন", "Role থেকে 'User' সিলেক্ট করুন এবং তার অফিস লোকেশন ও শিফট সিলেক্ট করে দিন।")
        add_num_step(doc, 5, "Save চাপুন", "'Save User' বাটনে চাপ দিন। সাথে সাথে তার আইডি তৈরি হয়ে যাবে এবং সে লগইন করতে পারবে।")
        
        # ৯. পাসওয়ার্ড রিসেট ও নিষ্ক্রিয় করা
        add_section_h2(doc, "৯. পাসওয়ার্ড রিসেট ও চাকরি ছাড়লে আইডি বন্ধ করা")
        add_bullet(doc, "পাসওয়ার্ড ভুলে গেলে", "ইউজার লিস্টে গিয়ে কর্মীর নামের ওপর চাপ দিন এবং 'Reset Password' বাটনে চাপ দিয়ে নতুন পাসওয়ার্ড লিখে দিন।")
        add_bullet(doc, "কর্মী চাকরি ছেড়ে দিলে", "ইউজার প্রোফাইলে গিয়ে 'Disable' বাটন চাপ দিন। সাথে সাথে তার আইডি বন্ধ হয়ে যাবে এবং সে আর অ্যাপে ঢুকতে পারবে না।")
        
        # ১০. অফিস লোকেশন ও শিফট সেট করা
        add_section_h2(doc, "১০. অফিস লোকেশন ও শিফট টাইম সেট করা")
        add_bullet(doc, "অফিস লোকেশন (Geofence)", "'Offices' মেনুতে গিয়ে আপনার অফিসের নাম লিখুন এবং ম্যাপে পয়েন্ট করে রেডিয়াস (যেমন ১০০ মিটার) সেট করে সেভ করুন। কর্মী এই দূরত্বের মধ্যে এসে হাজিরা দিলে 'Office Attendance' হিসেবে কাউন্ট হবে।")
        add_bullet(doc, "শিফট ও গ্রেস টাইম (Shifts)", "'Shifts' মেনুতে গিয়ে অফিসের সময় (যেমন সকাল ৯:০০ AM থেকে সন্ধ্যা ৬:০০ PM) এবং কত মিনিট পর্যন্ত লেট ছাড় দেওয়া হবে (যেমন ১৫ মিনিট) তা সেট করুন।")
        
        doc.add_page_break()
    
    # -------------------------------------------------------------
    # ENGLISH SECTION (ACTION ORIENTED)
    # -------------------------------------------------------------
    if include_en:
        add_section_h1(doc, "PART 2: ADMIN OPERATIONAL GUIDE (ENGLISH)")
        
        # 1. Admin Login
        add_section_h2(doc, "1. How to Login as Administrator")
        add_num_step(doc, 1, "Open App", "Launch the Smart Work Force application on your phone or tablet.")
        add_num_step(doc, 2, "Enter Admin Credentials", "Type your assigned Admin Username and Password.")
        add_num_step(doc, 3, "Tap LOGIN", "Tap 'LOGIN'. You will be taken straight to the Executive Dashboard.")
        add_num_step(doc, 4, "Toggle Language", "Tap `🌐 English / বাংলা` at the top right to switch language anytime.")
        
        # 2. Executive Dashboard
        add_section_h2(doc, "2. Executive Overview Dashboard")
        add_p(doc, "The home screen gives you an instant bird's-eye view of your entire organization:")
        add_bullet(doc, "Total Staff", "Total registered workforce in the system.")
        add_bullet(doc, "Active Now", "Number of field officers currently broadcasting live GPS.")
        add_bullet(doc, "Today's Present", "Employees who completed morning attendance.")
        add_bullet(doc, "Today's Absent", "Employees who have not punched in today.")
        add_bullet(doc, "Late Arrivals", "Staff who arrived after the shift grace time.")
        add_bullet(doc, "Attention Required Alert", "Instantly highlights if an employee turns off GPS or goes offline for more than 3 minutes. Tap the 'Call' button to phone them immediately.")
        
        # 3. Live Map Tracking
        add_section_h2(doc, "3. Real-Time Live Map Tracking")
        add_num_step(doc, 1, "Open Map", "Tap 'Live Map' from the bottom navigation menu.")
        add_num_step(doc, 2, "Inspect Marker Colors", "View all field staff pins on Google Maps:")
        add_bullet(doc, "🟢 Green Pin", "Officer is moving (traveling on bike/car; speed > 2 km/h).")
        add_bullet(doc, "🔵 Blue Pin", "Officer is stationary (at a client meeting or shop).")
        add_bullet(doc, "⚪ Gray Pin", "Offline (device disconnected or network lost for > 3 minutes).")
        add_bullet(doc, "🔴 Red Pin", "Officer manually disabled GPS on their phone.")
        add_num_step(doc, 3, "Tap Any Marker", "Opens the info card showing the officer's name, phone, battery %, current speed, and street address.")
        add_num_step(doc, 4, "Direct Phone Call", "Tap 'Call' to dial the employee directly.")
        
        # 4. Route History
        add_section_h2(doc, "4. Viewing Route History & Movement Playback")
        add_num_step(doc, 1, "Go to Route History", "Tap 'Route History' from the navigation drawer.")
        add_num_step(doc, 2, "Select Officer & Date", "Pick the employee's name and choose the date from the calendar.")
        add_num_step(doc, 3, "Review Daily Summary", "View total distance traveled (km), first punch-in time, and last punch-out time.")
        add_num_step(doc, 4, "Check Stops", "Numbered stop pins show every location where the officer stopped for over 5 minutes. Tap any pin to see arrival time, departure time, and duration.")
        add_num_step(doc, 5, "Simulate Playback", "Tap 'Play' to watch a video-like simulation of the officer traveling along the road.")
        
        # 5. Reviewing Attendance
        add_section_h2(doc, "5. Reviewing Attendance & Selfie Photos")
        add_num_step(doc, 1, "Go to Attendance", "Tap 'Attendance' from the menu.")
        add_num_step(doc, 2, "Browse Roster", "View attendance list for today or any past date.")
        add_num_step(doc, 3, "Inspect Selfie Photo", "Tap on any employee's photo thumbnail to view their full-resolution face photo.")
        add_num_step(doc, 4, "Verify Geofence", "Check whether the punch was recorded 'Inside Office Geofence' or 'Outdoor Field Duty'.")
        
        # 6. Customer Visits
        add_section_h2(doc, "6. Monitoring Customer Visits")
        add_num_step(doc, 1, "Go to Customer Visits", "Tap 'Customer Visits' from the menu.")
        add_num_step(doc, 2, "Inspect Reports", "View shops visited, meeting outcome, discussion notes, and shop storefront photos taken by the officer.")
        
        # 7. Approving Leaves
        add_section_h2(doc, "7. Approving or Rejecting Leave Requests")
        add_num_step(doc, 1, "Go to Leave", "Tap 'Leave' from the menu.")
        add_num_step(doc, 2, "Review Requests", "Check the applicant's name, leave type, requested dates, and stated reason.")
        add_num_step(doc, 3, "Approve or Reject", "Tap 'Approve' to grant the leave or 'Reject' to deny it. The employee receives an immediate notification.")
        
        # 8. User Management
        add_section_h2(doc, "8. Creating & Managing Employee Accounts")
        add_num_step(doc, 1, "Go to Users", "Tap 'User Management' from the menu.")
        add_num_step(doc, 2, "Tap '+ Create User'", "Tap the '+ Create User' button.")
        add_num_step(doc, 3, "Fill Details", "Enter Full Name, Phone Number, Username, and Password.")
        add_num_step(doc, 4, "Assign Role & Office", "Select Role ('User' or 'Admin'), assign their office location and duty shift, then tap 'Save User'.")
        add_bullet(doc, "Reset Password", "Tap on any user card and select 'Reset Password' to set a new password.")
        add_bullet(doc, "Disable Account", "If an employee leaves the company, tap 'Disable' to instantly block their access.")
        
        # 9. Office Geofences & Shifts
        add_section_h2(doc, "9. Office Geofence & Shift Setup")
        add_bullet(doc, "Office Locations", "Under 'Offices', add office branches with latitude, longitude, and allowed geofence radius (e.g. 50m to 100m).")
        add_bullet(doc, "Shifts", "Under 'Shifts', define shift start time (e.g. 09:00 AM), end time, and late grace period (e.g. 15 minutes).")

    doc.save(output_path)
    print(f"Saved Admin Manual: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\Shuvo\zynexbd\live_tracking"
    doc_dir = r"d:\Shuvo\zynexbd\live_tracking\LiveTrackingSystem\Documentation"
    
    # 1. Master Easy Admin Manual (BN + EN)
    master_path = os.path.join(out_dir, "Smart_WorkForce_Admin_Guideline.docx")
    build_pure_admin_manual(master_path, include_en=True, include_bn=True)
    
    doc_path = os.path.join(doc_dir, "Smart_WorkForce_Admin_Guideline.docx")
    build_pure_admin_manual(doc_path, include_en=True, include_bn=True)
    
    # 2. Standalone Bangla
    bn_path = os.path.join(out_dir, "Smart_WorkForce_Admin_Guideline_Bangla.docx")
    build_pure_admin_manual(bn_path, include_en=False, include_bn=True)
    
    # 3. Standalone English
    en_path = os.path.join(out_dir, "Smart_WorkForce_Admin_Guideline_English.docx")
    build_pure_admin_manual(en_path, include_en=True, include_bn=False)
    
    print("Pure admin manual docx generated!")
