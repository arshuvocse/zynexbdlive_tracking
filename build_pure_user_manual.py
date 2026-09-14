import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from admin_doc_builder import (
    set_cell_background, set_cell_margins, add_callout,
    add_styled_table, add_header_footer, add_section_h1,
    add_section_h2, add_section_h3, add_p, add_bullet, add_num_step
)

def build_pure_user_manual(output_path, include_en=True, include_bn=True):
    doc = Document()
    add_header_footer(doc, "Field User App Manual / কীভাবে অ্যাপ ব্যবহার করবেন")
    
    # Title Cover Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(18)
    title_p.paragraph_format.space_after = Pt(2)
    r_brand = title_p.add_run("SMART WORK FORCE")
    r_brand.bold = True
    r_brand.font.size = Pt(24)
    r_brand.font.name = "Segoe UI"
    r_brand.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) # Teal
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    r_sub = sub_p.add_run("User App Manual: How to Use the App (Step-by-Step)\nইউজার অ্যাপ ব্যবহারের নিয়ম (ধাপে ধাপে সহজ গাইড)")
    r_sub.font.size = Pt(13)
    r_sub.font.name = "Segoe UI"
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    add_callout(
        doc,
        "এই ম্যানুয়ালে কোনো কঠিন বা টেকনিক্যাল কথা নেই। আপনি কীভাবে অ্যাপে লগইন করবেন, কীভাবে সেলফি তুলে হাজিরা দেবেন, কীভাবে নতুন কাস্টমার অ্যাড করবেন, কাস্টমারকে কীভাবে কল দেবেন এবং কীভাবে ভিজিট রিপোর্ট জমা দেবেন — তা একদম সহজ ভাষায় বুঝিয়ে দেওয়া হয়েছে।\n"
        "This is a pure 'How-To' guide. It explains exactly what buttons to tap, what to fill in, and how to complete your daily work in the app without any technical jargon.",
        title="HOW TO USE THIS APP / অ্যাপ ব্যবহারের সহজ নিয়ম",
        border_hex="0F766E",
        bg_hex="F0FDF4"
    )
    
    # -------------------------------------------------------------
    # BANGLA SECTION (FIRST & PRIMARY)
    # -------------------------------------------------------------
    if include_bn:
        add_section_h1(doc, "প্রথম অংশ: অ্যাপ ব্যবহারের নিয়ম (বাংলা নির্দেশিকা)")
        
        # ১. লগইন করা
        add_section_h2(doc, "১. অ্যাপে লগইন করার নিয়ম")
        add_num_step(doc, 1, "অ্যাপ ওপেন করুন", "আপনার মোবাইলে Smart Work Force অ্যাপটি ওপেন করুন।")
        add_num_step(doc, 2, "ইউজারনেম ও পাসওয়ার্ড দিন", "Username এর ঘরে আপনার আইডি এবং Password এর ঘরে আপনার পাসওয়ার্ড লিখুন।")
        add_num_step(doc, 3, "লগইন চাপুন", "'LOGIN' বাটনে চাপ দিন।")
        add_num_step(doc, 4, "ভাষা নির্বাচন", "অ্যাপটি বাংলায় দেখতে চাইলে ওপরে ডানপাশে '🌐 বাংলা' বাটনে চাপ দিন।")
        
        # ২. প্রথমবার ফোনের পারমিশন অন করা
        add_section_h2(doc, "২. প্রথমবার অ্যাপ চালু করলে কী করবেন (পারমিশন)")
        add_p(doc, "অ্যাপটি ইনস্টল করার পর প্রথমবার চালু করলে মোবাইল থেকে ৪টি পারমিশন চাইবে। এগুলো একবার ঠিক করে নিলেই সারাদিন কোনো সমস্যা হবে না:")
        add_bullet(doc, "📍 লোকেশন (Location)", "'Allow all the time' অপশনে চাপ দিন। এতে মোবাইল পকেটে বা স্ক্রিন বন্ধ থাকলেও ট্র্যাকিং সচল থাকবে।")
        add_bullet(doc, "📷 ক্যামেরা (Camera)", "'While using the app' বা 'Allow' চাপুন (সেলফি হাজিরা ও ছবি তোলার জন্য)।")
        add_bullet(doc, "🔔 নোটিফিকেশন (Notification)", "'Allow' চাপুন (অফিসের নোটিশ পাওয়ার জন্য)।")
        add_bullet(doc, "🔋 ব্যাটারি সেভার (খুবই জরুরি)", "মোবাইলের Settings -> Apps -> Smart Work Force -> Battery তে গিয়ে 'No Restrictions' বা 'Unrestricted' সিলেক্ট করে দিন। (যাতে ফোন নিজ থেকে অ্যাপ বন্ধ না করে)।")
        
        # ৩. সকালে হাজিরা দেওয়া
        add_section_h2(doc, "৩. সকালে হাজিরা দেওয়া (Duty In - Punch In)")
        add_p(doc, "সকালে ডিউটি শুরুর সময় যা করবেন (সময় লাগবে মাত্র ৩০ সেকেন্ড):")
        add_num_step(doc, 1, "Duty In বাটনে চাপুন", "হোম স্ক্রিনের বড় সবুজ 'Duty In' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "সেলফি ছবি তুলুন", "সামনের ক্যামেরা অন হবে। ভালো আলোতে ক্যামেরার দিকে তাকিয়ে পরিষ্কার একটি সেলফি ছবি তুলুন।")
        add_num_step(doc, 3, "Confirm চাপুন", "ছবি ঠিক থাকলে নিচে 'Confirm' বাটনে চাপ দিন।")
        add_p(doc, "✅ সাথে সাথে আপনার উপস্থিতি রেকর্ড হয়ে যাবে, স্ক্রিনে ডিউটি টাইমার চালু হবে এবং আপনার কাজ শুরু হবে।")
        
        # ৪. নতুন কাস্টমার যোগ করা
        add_section_h2(doc, "৪. নতুন কাস্টমার বা দোকান অ্যাড করা")
        add_p(doc, "মাঠে গিয়ে কোনো নতুন দোকান বা ক্লায়েন্ট পেলে কীভাবে সেভ করবেন (সময় লাগবে ১ মিনিট):")
        add_num_step(doc, 1, "Customers এ যান", "স্ক্রিনের নিচের মেনু থেকে 'Customers' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "+ বাটন চাপুন", "উপরে ডানে থাকা '+' বা 'Add Customer' বাটনে চাপ দিন।")
        add_num_step(doc, 3, "নাম ও মোবাইল লিখুন", "দোকানের নাম এবং দোকানদারের মোবাইল নম্বর লিখুন।")
        add_num_step(doc, 4, "ঠিকানা", "আপনি যেখানে দাঁড়িয়ে আছেন সেখানকার ঠিকানা ও জিপিএস স্বয়ংক্রিয়ভাবে বসে যাবে, আপনাকে কিছু লিখতে হবে না।")
        add_num_step(doc, 5, "ছবি তুলুন (ঐচ্ছিক)", "ক্যামেরা আইকনে চাপ দিয়ে দোকানের সাইনবোর্ডের একটি ছবি তুলতে পারেন।")
        add_num_step(doc, 6, "Save Customer চাপুন", "নিচে 'Save Customer' বাটনে চাপ দিন। কাস্টমারটি সেভ হয়ে যাবে।")
        
        # ৫. কাস্টমারকে কল ও রাস্তা দেখা
        add_section_h2(doc, "৫. কাস্টমারকে সরাসরি কল ও ম্যাপে রাস্তা দেখা")
        add_p(doc, "কাস্টমার লিস্টে গিয়ে যেকোনো কাস্টমারের নামের পাশে থাকা বাটন চাপুন:")
        add_bullet(doc, "📞 ফোন করতে চাইলে", "ফোন আইকনে চাপ দিন। নম্বর ডায়াল না করেই সরাসরি আপনার ফোন থেকে কল চলে যাবে।")
        add_bullet(doc, "🧭 দোকানে যাওয়ার রাস্তা দেখতে", "ম্যাপ আইকনে চাপ দিন। গুগল ম্যাপ ওপেন হয়ে দোকান পর্যন্ত যাওয়ার সহজ রাস্তা দেখিয়ে দেবে।")
        
        # ৬. ভিজিট রিপোর্ট সাবমিট করা
        add_section_h2(doc, "৬. কাস্টমার ভিজিট এন্ট্রি দেওয়া (Record Visit)")
        add_p(doc, "দোকানে গিয়ে কথা বলার পর কীভাবে ভিজিট জমা দেবেন (সময় লাগবে ১ মিনিট):")
        add_num_step(doc, 1, "Record Visit চাপুন", "কাস্টমারের পাশে থাকা 'Record Visit' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "স্ট্যাটাস সিলেক্ট করুন", "ড্রপডাউন থেকে সিলেক্ট করুন কী হলো — 'আগ্রহী' / 'অর্ডার হয়েছে' / 'পেমেন্ট পেয়েছি' / 'আবার যেতে হবে'।")
        add_num_step(doc, 3, "নোট লিখুন", "Remarks এর ঘরে সংক্ষেপে ১ লাইনে লিখুন (যেমন: 'আগামী সোমবার মাল ডেলিভারি দিতে হবে')।")
        add_num_step(doc, 4, "পরবর্তী তারিখ দিন", "যদি আবার যেতে হয়, তবে কবে যাবেন সেই তারিখ ও সময় সিলেক্ট করুন।")
        add_num_step(doc, 5, "Submit Visit চাপুন", "'Submit Visit' বাটনে চাপ দিন। সাথে সাথে আপনার ভিজিট অফিসে জমা হয়ে যাবে।")
        
        # ৭. ফলো-আপ মিটিং চেক করা
        add_section_h2(doc, "৭. ফলো-আপ চেক করা (আজকে কার সাথে মিটিং)")
        add_num_step(doc, 1, "Follow-ups এ যান", "স্ক্রিনের নিচের মেনু থেকে 'Follow-ups' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "Today ট্যাব দেখুন", "আজকের দিনে যাদের সাথে দেখা করার কথা তাদের তালিকা দেখতে পাবেন।")
        add_num_step(doc, 3, "কল করুন বা ভিজিট দিন", "সরাসরি সেখান থেকে কাস্টমারকে কল দিতে পারবেন বা দেখা করতে পারবেন।")
        add_num_step(doc, 4, "কাজ শেষ হলে", "মিটিং শেষ হলে 'Mark Completed' বাটনে চাপ দিলে তালিকা থেকে সরে যাবে।")
        
        # ৮. ছুটির আবেদন করা
        add_section_h2(doc, "৮. মোবাইল থেকে ছুটির আবেদন করার নিয়ম")
        add_num_step(doc, 1, "Apply Leave চাপুন", "হোম স্ক্রিনের 'Apply Leave' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "ছুটির ধরন বাছুন", "নৈমিত্তিক ছুটি (Casual), অসুস্থতাজনিত ছুটি (Sick) নাকি বাৎসরিক ছুটি সিলেক্ট করুন।")
        add_num_step(doc, 3, "তারিখ সিলেক্ট করুন", "কবে থেকে কবে ছুটি চান ক্যালেন্ডার থেকে সিলেক্ট করুন।")
        add_num_step(doc, 4, "কারণ লিখে Submit চাপুন", "ছুটির কারণ লিখে 'Submit' বাটনে চাপ দিন। ম্যানেজার অনুমোদন করলে আপনার ফোনে নোটিফিকেশন আসবে।")
        
        # ৯. নিজের হাজিরা দেখা
        add_section_h2(doc, "৯. নিজের হাজিরা ও কাজের সময় দেখার নিয়ম")
        add_p(doc, "হোম স্ক্রিনের **'Attendance History'** বাটনে চাপ দিন:")
        add_bullet(doc, "সবুজ দিন", "উপস্থিত (Present)")
        add_bullet(doc, "হলুদ দিন", "দেরিতে হাজিরা (Late)")
        add_bullet(doc, "লাল দিন", "অনুপস্থিত (Absent)")
        add_p(doc, "যেকোনো দিনের ওপর চাপ দিলে সকালে কয়টায় পাঞ্চ করেছেন এবং বিকালে কয়টায় বের হয়েছেন তা পরিষ্কার দেখতে পাবেন।")
        
        # ১০. ডিউটি শেষে বের হওয়া
        add_section_h2(doc, "১০. ডিউটি শেষে বের হওয়া (Duty Out - Punch Out)")
        add_p(doc, "দিনের কাজ শেষ করার সময় যা করবেন:")
        add_num_step(doc, 1, "Duty Out বাটনে চাপুন", "হোম স্ক্রিনের বড় লাল 'Duty Out' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "সেলফি তুলুন", "একটি বিদায়ী সেলফি ছবি তুলুন।")
        add_num_step(doc, 3, "Confirm চাপুন", "'Confirm' বাটনে চাপ দিন।")
        add_p(doc, "✅ সারাদিনে আপনি কত ঘণ্টা কত মিনিট কাজ করেছেন তা স্ক্রিনে দেখতে পাবেন এবং ডিউটি শেষ হবে।")
        
        # জরুরি টিপস
        add_section_h2(doc, "১১. জরুরি ৩টি টিপস")
        add_bullet(doc, "১. ইন্টারনেট না থাকলে?", "কোনো চিন্তা নেই! নেট না থাকলেও আপনি হাজিরা ও ভিজিট দিতে পারবেন। ফোনে ইন্টারনেট আসবা মাত্রই অটোমেটিক অফিসে চলে যাবে।")
        add_bullet(doc, "২. স্ক্রিনের ওপরের লাইট", "স্ক্রিনের ওপরে 'GPS: চালু' এবং 'অনলাইন' লেখা সবুজ থাকলে বুঝবেন সব ঠিক আছে।")
        add_bullet(doc, "৩. ফোনের চার্জ বাঁচানো", "অপ্রয়োজনে বারবার ফোন রিস্টার্ট বা অ্যাপ বন্ধ (Force Stop) করবেন না।")
        
        doc.add_page_break()
    
    # -------------------------------------------------------------
    # ENGLISH SECTION (PURE ACTION GUIDE)
    # -------------------------------------------------------------
    if include_en:
        add_section_h1(doc, "PART 2: USER APP OPERATIONAL GUIDE (ENGLISH)")
        
        # 1. Login
        add_section_h2(doc, "1. How to Login")
        add_num_step(doc, 1, "Open the App", "Launch the Smart Work Force app on your mobile phone.")
        add_num_step(doc, 2, "Enter Username & Password", "Type your assigned Username and Password.")
        add_num_step(doc, 3, "Tap LOGIN", "Tap the 'LOGIN' button.")
        add_num_step(doc, 4, "Switch Language", "Tap `🌐 English / বাংলা` at the top right to toggle language at any time.")
        
        # 2. Permissions
        add_section_h2(doc, "2. One-Time Setup (Phone Permissions)")
        add_p(doc, "When opening the app for the first time, grant these 4 permissions so tracking runs smoothly all day:")
        add_bullet(doc, "📍 Location", "Select **'Allow all the time'** (so tracking works even when screen is locked in your pocket).")
        add_bullet(doc, "📷 Camera", "Select **'Allow'** (to take selfies for attendance and snap shop photos).")
        add_bullet(doc, "🔔 Notification", "Select **'Allow'** (to receive alerts and keep tracking service running).")
        add_bullet(doc, "🔋 Battery Saver", "Go to phone `Settings -> Apps -> Smart Work Force -> Battery` and choose **'No Restrictions'** or **'Unrestricted'**.")
        
        # 3. Morning Duty In
        add_section_h2(doc, "3. Morning Attendance (Duty In - Punch In)")
        add_p(doc, "Takes only 30 seconds every morning:")
        add_num_step(doc, 1, "Tap 'Duty In'", "On the home screen, tap the large green **'Duty In'** button.")
        add_num_step(doc, 2, "Snap a Selfie", "Position your face clearly in the camera viewfinder and take a selfie.")
        add_num_step(doc, 3, "Tap Confirm", "Tap **'Confirm'**. Your attendance is recorded and your workday timer begins!")
        
        # 4. Adding a Customer
        add_section_h2(doc, "4. How to Add a New Customer / Dealer")
        add_p(doc, "Takes less than 1 minute:")
        add_num_step(doc, 1, "Go to Customers", "Tap **'Customers'** on the bottom navigation bar.")
        add_num_step(doc, 2, "Tap '+ Add Customer'", "Tap the **'+'** or **'Add Customer'** button at the top.")
        add_num_step(doc, 3, "Enter Name & Mobile", "Type Customer/Shop Name and Phone Number.")
        add_num_step(doc, 4, "Address & GPS", "The app automatically captures your current GPS coordinates and street address.")
        add_num_step(doc, 5, "Take Shop Photo (Optional)", "Tap the camera icon to snap a photo of the shop signboard.")
        add_num_step(doc, 6, "Tap Save", "Tap **'Save Customer'**. The client is saved immediately.")
        
        # 5. Call & Navigate
        add_section_h2(doc, "5. How to Call a Customer or Open in Google Maps")
        add_p(doc, "Find any customer in the list and tap the action icons:")
        add_bullet(doc, "📞 Direct Phone Call", "Tap the phone icon to dial the customer instantly without memorizing numbers.")
        add_bullet(doc, "🧭 Driving Directions", "Tap the map icon to launch Google Maps navigation straight to their shop.")
        
        # 6. Record Visit
        add_section_h2(doc, "6. How to Submit a Customer Visit (Record Visit)")
        add_p(doc, "After meeting with a client:")
        add_num_step(doc, 1, "Tap 'Record Visit'", "Tap the visit button on the customer's card.")
        add_num_step(doc, 2, "Select Visit Status", "Choose from: Interested / Order Placed / Quotation Provided / Payment Received / Follow-up Needed.")
        add_num_step(doc, 3, "Write Short Note", "Type a 1-line note of what was agreed (e.g. 'Deliver 5 cartons next Monday').")
        add_num_step(doc, 4, "Pick Next Date", "If you need to return, select the date and time.")
        add_num_step(doc, 5, "Tap Submit Visit", "Tap **'Submit Visit'**. The report is synced to office immediately.")
        
        # 7. Follow-ups
        add_section_h2(doc, "7. How to Check Your Follow-up Schedule")
        add_num_step(doc, 1, "Go to Follow-ups", "Tap **'Follow-ups'** on the bottom bar.")
        add_num_step(doc, 2, "Check 'Today'", "See all clients you are scheduled to meet or call today.")
        add_num_step(doc, 3, "Call or Visit", "Tap to call directly or log the meeting.")
        add_num_step(doc, 4, "Mark Completed", "Tap **'Mark Completed'** when the task is done.")
        
        # 8. Apply Leave
        add_section_h2(doc, "8. How to Apply for Leave")
        add_num_step(doc, 1, "Tap 'Apply Leave'", "Tap **'Apply Leave'** on the home screen.")
        add_num_step(doc, 2, "Select Leave Type", "Choose Casual, Sick, or Annual leave.")
        add_num_step(doc, 3, "Pick Date Range", "Select Start Date and End Date.")
        add_num_step(doc, 4, "Enter Reason & Submit", "Write your reason and tap **'Submit'**. You will get an app notification when approved.")
        
        # 9. Attendance History
        add_section_h2(doc, "9. How to View Your Attendance & Working Hours")
        add_p(doc, "Tap **'Attendance History'** on the home screen to view your monthly calendar:")
        add_bullet(doc, "Green", "Present")
        add_bullet(doc, "Yellow", "Late Arrival")
        add_bullet(doc, "Red", "Absent")
        add_p(doc, "Tap any day to see your exact punch-in time, punch-out time, and total hours worked.")
        
        # 10. Duty Out
        add_section_h2(doc, "10. Evening Duty Out (Punch Out)")
        add_num_step(doc, 1, "Tap 'Duty Out'", "At the end of your shift, tap the large red **'Duty Out'** button.")
        add_num_step(doc, 2, "Take Exit Selfie", "Snap your closing selfie.")
        add_num_step(doc, 3, "Tap Confirm", "Tap **'Confirm'**. The screen shows your total work hours for the day.")
        
        # Important Tips
        add_section_h2(doc, "11. Important Field Tips")
        add_bullet(doc, "No Internet / Mobile Data?", "The app works 100% offline! It saves your attendance and visits and automatically uploads them the moment internet returns.")
        add_bullet(doc, "Top Screen Indicators", "Keep an eye on the top green pills: 'GPS: ON' and 'ONLINE'.")

    doc.save(output_path)
    print(f"Saved: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\Shuvo\zynexbd\live_tracking"
    doc_dir = r"d:\Shuvo\zynexbd\live_tracking\LiveTrackingSystem\Documentation"
    
    # 1. Master Easy User Manual (BN + EN)
    master_path = os.path.join(out_dir, "Smart_WorkForce_User_Guideline.docx")
    build_pure_user_manual(master_path, include_en=True, include_bn=True)
    
    doc_path = os.path.join(doc_dir, "Smart_WorkForce_User_Guideline.docx")
    build_pure_user_manual(doc_path, include_en=True, include_bn=True)
    
    # 2. Standalone Bangla
    bn_path = os.path.join(out_dir, "Smart_WorkForce_User_Guideline_Bangla.docx")
    build_pure_user_manual(bn_path, include_en=False, include_bn=True)
    
    # 3. Standalone English
    en_path = os.path.join(out_dir, "Smart_WorkForce_User_Guideline_English.docx")
    build_pure_user_manual(en_path, include_en=True, include_bn=False)
    
    print("Pure user manual docx generated!")
