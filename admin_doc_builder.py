import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text, title="IMPORTANT NOTICE / জরুরি নির্দেশনা", border_hex="1E3A8A", bg_hex="F0F7FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    c = tbl.cell(0, 0)
    c.width = Inches(6.5)
    set_cell_background(c, bg_hex)
    set_cell_margins(c, top=140, bottom=140, left=200, right=150)
    
    tcPr = c._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"📌 {title}\n")
    r_title.bold = True
    r_title.font.size = Pt(10.5)
    r_title.font.name = "Segoe UI"
    r_title.font.color.rgb = RGBColor.from_string(border_hex)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.name = "Segoe UI"
    r_text.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(0)
    p_end.paragraph_format.space_after = Pt(4)

def add_styled_table(doc, col_widths, headers, rows_data, header_bg="1E3A8A"):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header
    hdr = tbl.rows[0]
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.width = Inches(col_widths[i])
        set_cell_background(c, header_bg)
        set_cell_margins(c, top=120, bottom=120, left=140, right=140)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.name = "Segoe UI"
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for r_idx, r_data in enumerate(rows_data):
        row = tbl.add_row()
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.size = Pt(9)
            r.font.name = "Segoe UI"
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(0)
    p_end.paragraph_format.space_after = Pt(4)

def add_header_footer(doc, title_text):
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
        # Header
        hdr = s.header
        hp = hdr.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"Smart Work Force | {title_text}")
        hrun.font.size = Pt(8.5)
        hrun.font.name = "Segoe UI"
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        # Footer
        ftr = s.footer
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Smart Work Force • Confidential Operational Document • Enterprise Edition")
        frun.font.size = Pt(8)
        frun.font.name = "Segoe UI"
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

def add_section_h1(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(15)
    r.font.name = "Segoe UI"
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Blue
    
    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(6)
        r_sub = p_sub.add_run(subtitle)
        r_sub.italic = True
        r_sub.font.size = Pt(10)
        r_sub.font.name = "Segoe UI"
        r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

def add_section_h2(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = "Segoe UI"
    r.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) # Teal

def add_section_h3(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.name = "Segoe UI"
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.name = "Segoe UI"
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    r_pre = p.add_run(bold_prefix + ": ")
    r_pre.bold = True
    r_pre.font.size = Pt(10)
    r_pre.font.name = "Segoe UI"
    r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    r_txt = p.add_run(text)
    r_txt.font.size = Pt(10)
    r_txt.font.name = "Segoe UI"
    r_txt.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def add_num_step(doc, step_num, bold_title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    r_num = p.add_run(f"{step_num}. ")
    r_num.bold = True
    r_num.font.size = Pt(10)
    r_num.font.name = "Segoe UI"
    r_num.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    r_title = p.add_run(f"{bold_title} — ")
    r_title.bold = True
    r_title.font.size = Pt(10)
    r_title.font.name = "Segoe UI"
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    r_txt = p.add_run(text)
    r_txt.font.size = Pt(10)
    r_txt.font.name = "Segoe UI"
    r_txt.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

print("Formatting library ready.")
