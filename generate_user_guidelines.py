import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from admin_doc_builder import (
    set_cell_background, set_cell_margins, add_callout,
    add_styled_table, add_header_footer, add_section_h1,
    add_section_h2, add_section_h3, add_p, add_bullet, add_num_step
)

def build_easy_user_document(output_path, include_en=True, include_bn=True):
    doc = Document()
    add_header_footer(doc, "Field User Easy Manual / সহজ ব্যবহার নির্দেশিকা")
    
    # Title Cover Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(18)
    title_p.paragraph_format.space_after = Pt(2)
    r_brand = title_p.add_run("SMART WORK FORCE")
    r_brand.bold = True
    r_brand.font.size = Pt(24)
    r_brand.font.name = "Segoe UI"
    r_brand.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) # Teal
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    r_sub = sub_p.add_run("Field Officer & Employee Simple User Manual\nসহজ ফিল্ড অফিসার ও ইউজার ব্যবহার নির্দেশিকা")
    r_sub.font.size = Pt(13)
    r_sub.font.name = "Segoe UI"
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    # Quick Overview Box
    add_callout(
        doc,
        "এই গাইডটি মাঠপর্যায়ের কর্মকর্তা, সেলস রিপ্রেজেন্টেটিভ এবং কর্মচারীদের জন্য অত্যন্ত সহজ ভাষায় তৈরি করা হয়েছে। কোনো জটিলতা ছাড়া প্রতিদিনের কাজ কীভাবে করবেন তা ধাপে ধাপে বুঝিয়ে দেওয়া হলো।\n"
        "This guide is written in plain, simple language for field officers and staff. Follow these simple numbered steps to complete your daily duty easily.",
        title="EASY START GUIDE / সহজ নির্দেশিকা",
        border_hex="0F766E",
        bg_hex="F0FDF4"
    )
    
    # -------------------------------------------------------------
    # BANGLA SECTION (FIRST - AS REQUESTED FOR MAXIMUM EASE)
    # -------------------------------------------------------------
    if include_bn:
        add_section_h1(doc, "প্রথম অংশ: সহজ ইউজার নির্দেশিকা (বাংলা সংস্করণ)")
        
        # ১. এক নজরে প্রতিদিনের ৪টি কাজ
        add_section_h2(doc, "১. এক নজরে প্রতিদিনের ৪টি কাজ (Daily Routine)")
        add_p(doc, "প্রতিদিন আপনার ডিউটি চলাকালীন শুধু নিচের ৪টি কাজ করবেন:")
        
        routine_headers = ["সময় / ধাপ", "আপনাকে কী করতে হবে", "সময় লাগবে"]
        routine_rows = [
            ["১. সকালের হাজিরা (Duty In)", "কর্মস্থলে এসে অ্যাপ খুলে সেলফি তুলে 'Duty In' চাপুন।", "৩০ সেকেন্ড"],
            ["২. কাস্টমার ভিজিট (Visit)", "দোকানে পৌঁছে 'Record Visit' চেপে কী কথা হলো লিখুন।", "১ মিনিট"],
            ["৩. ফলো-আপ চেক (Follow-up)", "আজকে কার কার সাথে দেখা করার কথা তা দেখে কল দিন।", "যেকোনো সময়"],
            ["৪. সন্ধ্যার ছুটি (Duty Out)", "কাজ শেষ হলে আবার সেলফি তুলে 'Duty Out' চাপুন।", "৩০ সেকেন্ড"]
        ]
        add_styled_table(doc, [2.0, 3.2, 1.3], routine_headers, routine_rows, header_bg="0F766E")
        
        # ২. প্রথমবার অ্যাপ চালু ও মোবাইল সেটিং
        add_section_h2(doc, "২. প্রথমবার ফোন সেটআপ (মাত্র ২ মিনিট)")
        add_p(doc, "অ্যাপটি ইনস্টল করার পর প্রথমবার চালু করলে ৪টি পারমিশন চাইবে। এগুলো সঠিকভাবে 'Allow' করা খুবই জরুরি:")
        
        add_num_step(doc, 1, "লোকেশন পারমিশন (Location)", "মোবাইলে পারমিশনের বক্স আসলে 'Allow all the time' অপশনটি বেছে নিন। (এতে মোবাইল পকেটে বা লক থাকলেও আপনার হাজিরা ও ট্র্যাকিং সঠিকভাবে চলবে)।")
        add_num_step(doc, 2, "ক্যামেরা পারমিশন (Camera)", "'While using the app' বা 'Allow' দিন (সেলফি হাজিরা এবং দোকানের ছবি তোলার জন্য)।")
        add_num_step(doc, 3, "নোটিফিকেশন পারমিশন (Notification)", "'Allow' দিন (অফিসের নোটিশ ও মেসেজ পাওয়ার জন্য)।")
        add_num_step(doc, 4, "ব্যাটারি সেটিং (সবচেয়ে জরুরি)", "ফোনের ব্যাটারি সেভার যাতে অ্যাপ বন্ধ না করে, সেজন্য 'No Restrictions' বা 'সীমাহীন' নির্বাচন করুন।")
        
        add_section_h3(doc, "বিভিন্ন ফোনে ব্যাটারি সেটিং ঠিক করার সহজ উপায়:")
        add_bullet(doc, "শাওমি / রেডমি (Xiaomi/Redmi)", "ফোনের Settings -> Apps -> Manage Apps -> Smart Work Force -> Battery Saver এ গিয়ে 'No Restrictions' সিলেক্ট করুন এবং 'Autostart' চালু করুন।")
        add_bullet(doc, "স্যামসাং (Samsung)", "Settings -> Apps -> Smart Work Force -> Battery তে গিয়ে 'Unrestricted' দিন।")
        add_bullet(doc, "ভিভো / অপো / রিয়েলমি (Vivo/Oppo/Realme)", "Settings -> Battery তে গিয়ে অ্যাপটির জন্য Background Activity বা High Background Power চালু করুন।")
        
        # ৩. লগইন ও হোম স্ক্রিন
        add_section_h2(doc, "৩. লগইন করা ও হোম স্ক্রিন বোঝা")
        add_num_step(doc, 1, "লগইন করুন", "অফিস থেকে পাওয়া আপনার ইউজারনেম এবং পাসওয়ার্ড দিয়ে 'LOGIN' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "ভাষা পরিবর্তন", "ওপরে থাকা '🌐 English / বাংলা' বাটনে চাপ দিয়ে যেকোনো সময় সম্পূর্ণ অ্যাপ বাংলায় দেখতে পারবেন।")
        add_num_step(doc, 3, "স্ক্রিনের ওপরের ৩টি সিগন্যাল দেখুন", "ড্যাশবোর্ডের শীর্ষে ৩টি রঙিন গোল বাটন থাকবে:")
        add_bullet(doc, "জিপিএস সিগন্যাল", "সবুজ রঙে 'জিপিএস: চালু' থাকলে ঠিক আছে। লাল থাকলে ফোনের লোকেশন অন করুন।")
        add_bullet(doc, "নেটওয়ার্ক সিগন্যাল", "সবুজ 'অনলাইন' থাকলে সার্ভারের সাথে যুক্ত আছে।")
        add_bullet(doc, "ব্যাটারি সিগন্যাল", "আপনার ফোনের বর্তমান চার্জের পরিমাণ দেখাবে।")
        
        # ৪. সেলফি হাজিরা (Duty In ও Duty Out)
        add_section_h2(doc, "৪. হাজিরা দেওয়ার সহজ নিয়ম (ডিউটি ইন ও আউট)")
        
        add_section_h3(doc, "সকালে ডিউটি ইন (Punch In) করার ধাপ:")
        add_num_step(doc, 1, "বাটনে চাপ দিন", "হোম স্ক্রিনের বড় সবুজ 'Duty In' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "সেলফি ছবি তুলুন", "ক্যামেরা ওপেন হলে ভালো আলোতে আপনার মুখমণ্ডল ফ্রেমের ভেতরে রেখে ছবি তুলুন।")
        add_num_step(doc, 3, "কনফার্ম করুন", "ছবি ঠিক থাকলে 'Confirm' বাটনে চাপ দিন। ব্যস! হাজিরা সম্পন্ন হলো এবং আপনার কাজের টাইমার শুরু হয়ে গেল।")
        
        add_section_h3(doc, "কাজ শেষে ডিউটি আউট (Punch Out) করার ধাপ:")
        add_num_step(doc, 1, "বাটনে চাপ দিন", "ডিউটি শেষে হোম স্ক্রিনের লাল 'Duty Out' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "সেলফি ও কনফার্ম", "একটি বিদায়ী সেলফি তুলে কনফার্ম করুন।")
        add_num_step(doc, 3, "কাজের হিসাব", "সারাদিনে আপনি কত ঘণ্টা কত মিনিট কাজ করেছেন তা স্ক্রিনে সাথে সাথে দেখতে পাবেন।")
        
        # ৫. কাস্টমার ও দোকান যোগ করা
        add_section_h2(doc, "৫. নতুন কাস্টমার বা দোকান যোগ করা (১ মিনিটে)")
        add_p(doc, "ফিল্ডে গিয়ে নতুন কোনো দোকান বা ক্লায়েন্ট পেলে সহজে সেভ করে রাখুন:")
        add_num_step(doc, 1, "Add Customer চাপুন", "নিচের মেনু থেকে Customers এ গিয়ে ওপরে থাকা '+' বা 'Add Customer' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "দোকানের নাম ও মোবাইল নম্বর লিখুন", "দোকানদারের নাম এবং মোবাইল নম্বর দিন।")
        add_num_step(doc, 3, "স্বয়ংক্রিয় ঠিকানা", "আপনি যেখানে দাঁড়িয়ে আছেন সেখানকার জিপিএস ও ঠিকানা অ্যাপ নিজে নিজেই নিয়ে নেবে।")
        add_num_step(doc, 4, "ছবি তুলুন (ঐচ্ছিক)", "প্রয়োজনে দোকানের সাইনবোর্ডের একটি ছবি তুলতে পারেন।")
        add_num_step(doc, 5, "সেভ করুন", "'Save Customer' বাটনে চাপ দিন। কাস্টমার যুক্ত হয়ে যাবে।")
        
        add_section_h3(doc, "কাস্টমারকে সরাসরি কল ও ম্যাপে রাস্তা দেখা:")
        add_bullet(doc, "সরাসরি ফোন কল", "কাস্টমারের পাশে থাকা 📞 ফোন আইকনে চাপ দিলে কোনো নম্বর না তুলেই সরাসরি কল চলে যাবে।")
        add_bullet(doc, "গুগল ম্যাপে রাস্তা", "🧭 ম্যাপ আইকনে চাপ দিলে গুগল ম্যাপে দোকানের রাস্তা এবং দূরত্ব দেখাবে।")
        
        # ৬. ফিল্ড ভিজিট রিপোর্ট দেওয়া
        add_section_h2(doc, "৬. ফিল্ড ভিজিট রিপোর্ট সাবমিট করা (১ মিনিট)")
        add_p(doc, "যেকোনো দোকানে ভিজিট শেষ করার পর রিপোর্ট এন্ট্রি করার নিয়ম:")
        add_num_step(doc, 1, "'Record Visit' চাপুন", "কাস্টমারের কার্ডে থাকা 'Record Visit' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "ফলাফল বেছে নিন", "ড্রপডাউন থেকে সিলেক্ট করুন (যেমন: আগ্রহী / অর্ডার হয়েছে / কোটেশন দিয়েছি / পেমেন্ট পেয়েছি / ফলো-আপ লাগবে)।")
        add_num_step(doc, 3, "ছোট নোট লিখুন", "কী কথা হলো সংক্ষেপে ১ লাইনে লিখুন (যেমন: 'আগামী সোমবার ৫ কার্টুন ডেলিভারি দিতে বলেছে')।")
        add_num_step(doc, 4, "পরবর্তী তারিখ (যদি লাগে)", "পরে আবার কবে যেতে হবে সেই তারিখ সিলেক্ট করুন।")
        add_num_step(doc, 5, "সাবমিট করুন", "'Submit Visit' চাপুন। সাথে সাথে আপনার ভিজিট অফিসে জমা হয়ে যাবে।")
        
        # ৭. ফলো-আপ রিমাইন্ডার
        add_section_h2(doc, "৭. ফলো-আপ তালিকা ও রিমাইন্ডার")
        add_p(doc, "নিচের মেনু থেকে 'Follow-ups' এ গেলে ৪টি ট্যাব দেখতে পাবেন:")
        add_bullet(doc, "আজকে (Today)", "আজকের দিনে যাদের সাথে দেখা বা কথা বলার কথা।")
        add_bullet(doc, "আগামীকাল (Tomorrow)", "আগামীকালের মিটিংয়ের তালিকা।")
        add_bullet(doc, "আসন্ন (Upcoming)", "সামনের দিনগুলোর শিডিউল।")
        add_bullet(doc, "বকেয়া (Overdue)", "যেসব কাজের সময় পার হয়ে গেছে কিন্তু এখনো শেষ করেননি।")
        add_p(doc, "যেকোনো কার্ড থেকে সরাসরি কাস্টমারকে কল করতে পারবেন অথবা কাজ শেষ হলে 'Mark Completed' বাটনে চাপ দিয়ে তালিকা ক্লিয়ার করতে পারবেন।")
        
        # ৮. ছুটির আবেদন
        add_section_h2(doc, "৮. মোবাইল থেকে ছুটির আবেদন করা")
        add_num_step(doc, 1, "'Apply Leave' চাপুন", "হোম স্ক্রিনে থাকা 'Apply Leave' বাটনে চাপ দিন।")
        add_num_step(doc, 2, "ছুটির ধরন বেছে নিন", "নৈমিত্তিক ছুটি (Casual), অসুস্থতাজনিত ছুটি (Sick) অথবা বাৎসরিক ছুটি সিলেক্ট করুন।")
        add_num_step(doc, 3, "তারিখ দিন", "কবে থেকে কবে ছুটি চান তা সিলেক্ট করুন (মোট দিন স্বয়ংক্রিয়ভাবে হিসাব হবে)।")
        add_num_step(doc, 4, "কারণ লিখে সাবমিট দিন", "ছুটির কারণ লিখে 'Submit' বাটনে চাপ দিন। আপনার সুপারভাইজার অনুমোদন করলে মোবাইলে নোটিফিকেশন আসবে।")
        
        # ৯. মাসিক হাজিরা চেক
        add_section_h2(doc, "৯. নিজের মাসিক হাজিরা ও কাজের রেকর্ড দেখা")
        add_p(doc, "হোম স্ক্রিন থেকে 'Attendance History' তে গেলে ক্যালেন্ডার দেখতে পাবেন:")
        add_bullet(doc, "সবুজ রঙ", "উপস্থিত (Present)")
        add_bullet(doc, "হলুদ রঙ", "দেরিতে হাজিরা (Late)")
        add_bullet(doc, "লাল রঙ", "অনুপস্থিত (Absent)")
        add_bullet(doc, "নীল রঙ", "সরকারি/সাপ্তাহিক ছুটি (Holiday)")
        add_p(doc, "যেকোনো দিনে চাপ দিলে কয়টায় ডিউটি ইন করেছেন, কয়টায় আউট করেছেন এবং মোট কত ঘণ্টা কাজ করেছেন তা দেখা যাবে।")
        
        # ১০. সাধারণ সমস্যা ও সহজ সমাধান
        add_section_h2(doc, "১০. সাধারণ সমস্যা ও সহজ সমাধান (FAQ)")
        add_bullet(doc, "ইন্টারনেট বা এমবি শেষ হয়ে গেলে কি কাজ করবে?", "হ্যাঁ, কাজ করবে! ইন্টারনেট না থাকলেও আপনি হাজিরা দিতে পারবেন এবং ভিজিট সেভ করতে পারবেন। পরে ফোনে ইন্টারনেট আসবা মাত্র তা নিজে নিজেই অফিসে জমা হয়ে যাবে।")
        add_bullet(doc, "জিপিএস লোকেশন না পেলে কী করবেন?", "ফোনের ওপর থেকে নোটিফিকেশন বার নামিয়ে Location অফ করে আবার অন করুন। বহুতল ভবনের নিচে থাকলে কয়েক সেকেন্ড খোলা আকাশের নিচে আসুন।")
        add_bullet(doc, "স্ক্রিন অফ বা লক করলে ট্র্যাকিং বন্ধ হয়ে যায় কেন?", "ফোনের ব্যাটারি সেভার অ্যাপ বন্ধ করে দিচ্ছে। ধাপ ২ অনুযায়ী ফোনের সেটিংসে গিয়ে Smart Work Force অ্যাপের ব্যাটারি 'No Restrictions' বা 'সীমাহীন' করে দিন।")
        
        doc.add_page_break()
    
    # -------------------------------------------------------------
    # ENGLISH SECTION (EASY & DIRECT)
    # -------------------------------------------------------------
    if include_en:
        add_section_h1(doc, "PART 2: FIELD OFFICER & EMPLOYEE EASY MANUAL (ENGLISH)")
        
        # 1. Daily 4-Step Routine
        add_section_h2(doc, "1. Daily 4-Step Work Routine")
        add_p(doc, "Your daily workday in Smart Work Force takes only four simple actions:")
        
        routine_en_headers = ["Step / Time", "What You Need to Do", "Duration"]
        routine_en_rows = [
            ["1. Morning Duty In", "Reach your work location, open app, snap selfie, tap 'Duty In'.", "30 Seconds"],
            ["2. Customer Visits", "Visit client/dealer, tap 'Record Visit', enter brief note & submit.", "1 Minute"],
            ["3. Follow-Ups", "Check today's scheduled meetings and call clients directly.", "Anytime"],
            ["4. Evening Duty Out", "End of shift: snap closing selfie, tap 'Duty Out', view total hours.", "30 Seconds"]
        ]
        add_styled_table(doc, [2.0, 3.2, 1.3], routine_en_headers, routine_en_rows, header_bg="0F766E")
        
        # 2. First-Time Phone Setup
        add_section_h2(doc, "2. First-Time Phone Setup (Only 2 Minutes)")
        add_p(doc, "When opening the app for the first time, grant these essential permissions so tracking works smoothly all day:")
        
        add_num_step(doc, 1, "Location Permission", "Select **'Allow all the time'**. This ensures tracking continues even when the phone is locked in your pocket.")
        add_num_step(doc, 2, "Camera Permission", "Select **'While using the app'** or **'Allow'** (needed for selfie attendance and shop photos).")
        add_num_step(doc, 3, "Notification Permission", "Select **'Allow'** (to receive admin alerts and keep tracking active).")
        add_num_step(doc, 4, "Battery Optimization (Most Important)", "Set battery mode to **'No Restrictions'** / Unrestricted so your phone doesn't kill the app.")
        
        add_section_h3(doc, "Brand-Specific Quick Battery Settings:")
        add_bullet(doc, "Xiaomi / Redmi / POCO", "Go to Settings -> Apps -> Manage Apps -> Smart Work Force -> Battery Saver -> Choose **'No Restrictions'** and turn ON **'Autostart'**.")
        add_bullet(doc, "Samsung Galaxy", "Go to Settings -> Apps -> Smart Work Force -> Battery -> Choose **'Unrestricted'**.")
        add_bullet(doc, "Vivo / OPPO / Realme", "Go to Settings -> Battery -> Enable Background Activity or High Background Power for the app.")
        
        # 3. Login & Status Indicators
        add_section_h2(doc, "3. Login & Top Screen Indicators")
        add_num_step(doc, 1, "Login", "Enter your assigned Username and Password, then tap **'LOGIN'**.")
        add_num_step(doc, 2, "Language Toggle", "Tap `🌐 English / বাংলা` at the top right to switch languages instantly.")
        add_num_step(doc, 3, "Check Status Indicators", "Observe the three colored status pills at the top:")
        add_bullet(doc, "GPS Pill", "Green **'GPS: ON'** means GPS is active. If red, turn on phone GPS.")
        add_bullet(doc, "Network Pill", "Green **'ONLINE'** means connected to server. Red means no internet.")
        add_bullet(doc, "Battery Pill", "Displays your current battery percentage.")
        
        # 4. Punch In & Punch Out
        add_section_h2(doc, "4. Punch In & Punch Out (Selfie Attendance)")
        
        add_section_h3(doc, "Morning Duty In (Punch In):")
        add_num_step(doc, 1, "Tap 'Duty In'", "Tap the big green **'Duty In'** button on the home screen.")
        add_num_step(doc, 2, "Snap Selfie", "Position your face inside the camera frame with good light and capture a clear selfie.")
        add_num_step(doc, 3, "Confirm", "Tap **'Confirm Photo'**. Your duty starts, the timer begins, and tracking starts automatically.")
        
        add_section_h3(doc, "Evening Duty Out (Punch Out):")
        add_num_step(doc, 1, "Tap 'Duty Out'", "At shift end, tap the red **'Duty Out'** button.")
        add_num_step(doc, 2, "Closing Selfie", "Capture your exit selfie and confirm.")
        add_num_step(doc, 3, "View Working Hours", "Your total work hours for the day will appear on screen.")
        
        # 5. Customer Directory & Adding New Customers
        add_section_h2(doc, "5. Customer Directory & Adding Customers")
        add_p(doc, "Keep all client details in your pocket:")
        add_num_step(doc, 1, "Add Customer", "Go to **'Customers'** from the bottom menu and tap **'+ Add Customer'**.")
        add_num_step(doc, 2, "Enter Name & Mobile", "Type the customer's shop name and phone number.")
        add_num_step(doc, 3, "Auto GPS Address", "Your current GPS location and street address are filled in automatically.")
        add_num_step(doc, 4, "Storefront Photo (Optional)", "Take a quick photo of the shop sign or business card.")
        add_num_step(doc, 5, "Save", "Tap **'Save Customer'**.")
        
        add_section_h3(doc, "1-Tap Direct Call & Navigation:")
        add_bullet(doc, "Direct Phone Call", "Tap the 📞 phone icon next to any customer to dial them instantly.")
        add_bullet(doc, "Google Maps Directions", "Tap the 🧭 navigation icon to view the driving route on Google Maps.")
        
        # 6. Logging Field Visits
        add_section_h2(doc, "6. Logging a Field Visit (Under 1 Minute)")
        add_p(doc, "When you finish meeting a dealer or customer:")
        add_num_step(doc, 1, "Tap 'Record Visit'", "Tap the visit button on the customer's card.")
        add_num_step(doc, 2, "Select Visit Status", "Choose outcome: Interested / Order Placed / Quotation Given / Payment Collected / Follow-up Needed.")
        add_num_step(doc, 3, "Type Brief Note", "Write a 1-line summary of what was discussed.")
        add_num_step(doc, 4, "Pick Next Follow-up Date", "If you need to meet again, pick the date and time.")
        add_num_step(doc, 5, "Submit", "Tap **'Submit Visit'**. Your visit is recorded with exact GPS proof.")
        
        # 7. Follow-Ups
        add_section_h2(doc, "7. Follow-up Reminders")
        add_p(doc, "Tap **'Follow-ups'** on the bottom navigation bar to view your tasks:")
        add_bullet(doc, "Today", "Customers scheduled for today.")
        add_bullet(doc, "Tomorrow", "Upcoming meetings scheduled for tomorrow.")
        add_bullet(doc, "Upcoming", "Future appointments.")
        add_bullet(doc, "Overdue", "Tasks that were missed or need attention.")
        add_p(doc, "Call clients directly from each card or tap **'Mark Completed'** when resolved.")
        
        # 8. Leave Applications
        add_section_h2(doc, "8. Applying for Leave")
        add_num_step(doc, 1, "Tap 'Apply Leave'", "Select **'Apply Leave'** from the home screen.")
        add_num_step(doc, 2, "Choose Leave Type", "Pick Casual, Sick, or Annual leave.")
        add_num_step(doc, 3, "Select Dates", "Choose Start Date and End Date.")
        add_num_step(doc, 4, "Type Reason & Submit", "Write your reason and tap **'Submit'**. You will get a notification when approved.")
        
        # 9. Attendance History
        add_section_h2(doc, "9. Viewing Your Monthly Duty Logs")
        add_p(doc, "Tap **'Attendance History'** to see your monthly calendar:")
        add_bullet(doc, "Green", "Present")
        add_bullet(doc, "Yellow", "Late Arrival")
        add_bullet(doc, "Red", "Absent")
        add_bullet(doc, "Blue", "Holiday / Weekend")
        add_p(doc, "Tap any day to see your exact Punch In and Punch Out times and total working hours.")
        
        # 10. FAQs & Quick Fixes
        add_section_h2(doc, "10. Common Problems & Quick Fixes (FAQs)")
        add_bullet(doc, "Does the app work without internet?", "Yes! The app saves all your visits and locations offline. As soon as your phone gets mobile data or Wi-Fi, it automatically uploads everything to the server.")
        add_bullet(doc, "GPS Location is inaccurate?", "Turn phone Location OFF and then ON again. If inside a basement, step outside for a few seconds.")
        add_bullet(doc, "Tracking stops when the screen is locked?", "Your phone's battery saver is closing the app. Re-visit Section 2 and set Smart Work Force battery to **'No Restrictions'** / Unrestricted.")

    doc.save(output_path)
    print(f"Successfully generated Easy User Manual: {output_path}")

if __name__ == "__main__":
    out_dir = r"d:\Shuvo\zynexbd\live_tracking"
    doc_dir = r"d:\Shuvo\zynexbd\live_tracking\LiveTrackingSystem\Documentation"
    
    # 1. Master Easy User Document (BN + EN)
    master_path = os.path.join(out_dir, "Smart_WorkForce_User_Guideline.docx")
    build_easy_user_document(master_path, include_en=True, include_bn=True)
    
    # Copy to Documentation folder
    doc_path = os.path.join(doc_dir, "Smart_WorkForce_User_Guideline.docx")
    build_easy_user_document(doc_path, include_en=True, include_bn=True)
    
    # 2. Standalone Bangla
    bn_path = os.path.join(out_dir, "Smart_WorkForce_User_Guideline_Bangla.docx")
    build_easy_user_document(bn_path, include_en=False, include_bn=True)
    
    # 3. Standalone English
    en_path = os.path.join(out_dir, "Smart_WorkForce_User_Guideline_English.docx")
    build_easy_user_document(en_path, include_en=True, include_bn=False)
    
    print("Easy User Guidelines generated successfully!")
